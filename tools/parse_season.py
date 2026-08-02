# -*- coding: utf-8 -*-
"""2026年 種付け順番表.xlsx(全日シート)を解析してJSONにする。

使い方:
  python tools/parse_season.py "<順番表.xlsxのパス>" tools/season2026.json

出力JSON: { "4-13": { "8:00": [entry...], "13:00": [...], "17:00": [...] }, ... }
entry = { mareName, sireCode, farm, kind, apptTime, note }
※出力JSONは実データ(外部譲渡禁止)なのでリポジトリにコミットしないこと(.gitignore済)
"""
import io
import json
import re
import sys
import unicodedata
import datetime

import openpyxl


def norm(s):
    return unicodedata.normalize("NFKC", str(s)) if s is not None else ""


def raw(s):
    return "" if s is None else str(s)


def fmt_time(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.time, datetime.datetime)):
        h, m = v.hour, v.minute
        if h == 0 and m == 0:
            return ""
        return f"{h}:{m:02d}"
    s = norm(v).strip()
    if s in ("0", "00:00:00", "0:00", ""):
        return ""
    return s.replace("：", ":")


FARM_SUFFIX = ("牧場", "Ｆ", "farm", "Farm", "F")
CODE_RE = r"[Ａ-Ｚａ-ｚA-Za-z]{2,4}"


def is_farm(s):
    s2 = s.strip()
    if not s2:
        return False
    if s2.startswith("（") or s2.startswith("("):
        return True
    return any(s2.endswith(suf) for suf in FARM_SUFFIX)


def strip_trailing_code(s):
    m = re.match(r"^(.*?)[\s　]*" + CODE_RE + r"より[\s　]*$", s)
    return m.group(1).strip() if m else s


def clean_farm(s):
    s = strip_trailing_code(s.strip())
    if (s.startswith("（") and s.endswith("）")) or (s.startswith("(") and s.endswith(")")):
        s = s[1:-1]
    return s.strip()


def clean_mare(s):
    return strip_trailing_code(s.strip()).strip()


def parse_sheet(ws):
    # 種牡馬は2行ペア(英語名行+カタカナ名行、C列コードが同一)
    pairs = []
    r = 4
    while r < 107:
        c1 = ws.cell(row=r, column=3).value
        c2 = ws.cell(row=r + 1, column=3).value
        if c1 and c2 and norm(c1) == norm(c2):
            pairs.append((r, r + 1))
            r += 2
        else:
            r += 1

    def get_note(top, bot):
        parts = []
        loc = norm(ws.cell(row=top, column=20).value).strip() or norm(
            ws.cell(row=bot, column=20).value
        ).strip()
        if loc:
            parts.append(loc)
        ov = norm(ws.cell(row=top, column=21).value).strip() or norm(
            ws.cell(row=bot, column=21).value
        ).strip()
        if ov == "OV":
            parts.append("OV")
        color = norm(ws.cell(row=top, column=26).value).strip() or norm(
            ws.cell(row=bot, column=26).value
        ).strip()
        if color == "赤":
            parts.append("上り")
        elif color in ("青", "靑"):
            parts.append("上り再発")
        return "　".join(parts)

    groups = {"8:00": [], "13:00": [], "17:00": []}
    for top, bot in pairs:
        code = raw(ws.cell(row=top, column=3).value).strip()
        note = get_note(top, bot)
        # E/I/M=予約時間, G/K/O=牧場or牝馬, H/L/P=新/再
        for label, tc, fc, kc in [("8:00", 5, 7, 8), ("13:00", 9, 11, 12), ("17:00", 13, 15, 16)]:
            t = fmt_time(ws.cell(row=top, column=tc).value)
            a = raw(ws.cell(row=top, column=fc).value)
            b = raw(ws.cell(row=bot, column=fc).value)
            kind = norm(ws.cell(row=bot, column=kc).value).strip() or norm(
                ws.cell(row=top, column=kc).value
            ).strip()
            af, bf = is_farm(norm(a)), is_farm(norm(b))
            if af and not bf:
                farm_raw, mare_raw = a, b
            elif bf and not af:
                farm_raw, mare_raw = b, a
            else:
                farm_raw, mare_raw = a, b
            if not mare_raw or norm(mare_raw).strip() in ("（ＮＦ）", "ー", "-", ""):
                continue
            mare = clean_mare(mare_raw)
            farm = clean_farm(farm_raw)
            if not mare:
                continue
            kindn = "新" if "新" in kind else ("再" if "再" in kind else "")
            groups[label].append(
                {
                    "mareName": mare,
                    "sireCode": code,
                    "farm": farm,
                    "kind": kindn,
                    "apptTime": t,
                    "note": note,
                }
            )
    return groups


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    xlsx, outp = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    season = {}
    for nm in wb.sheetnames:
        if not re.match(r"^\d+-\d+$", nm):
            continue
        g = parse_sheet(wb[nm])
        if sum(len(v) for v in g.values()) == 0:
            continue  # 空シート(テンプレ)は飛ばす
        season[nm] = g
    io.open(outp, "w", encoding="utf-8").write(json.dumps(season, ensure_ascii=False))
    total = sum(len(v) for g in season.values() for v in g.values())
    print(f"days={len(season)} entries={total} -> {outp}")


if __name__ == "__main__":
    main()
